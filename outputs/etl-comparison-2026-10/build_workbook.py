"""Build Takasbank-ETL-Comparison.xlsx from comparison_data.py.

Usage:
    python build_workbook.py          # writes the workbook
    python build_workbook.py --json   # also writes comparison-data.json for the canvas
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import comparison_data as d

HERE = Path(__file__).parent
OUT = HERE / "Takasbank-ETL-Comparison.xlsx"

NAVY = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="DCE6F1")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
COND_FILL = PatternFill("solid", fgColor="FFEB9C")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16, color=NAVY)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

TOOL_IDS = [t["id"] for t in d.TOOLS]
TOOL_SHORT = [t["short"] for t in d.TOOLS]


def header(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.fill = HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = CENTER
        c.border = BORDER


def body(ws, row, values, start_col=1, wrap=True):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.border = BORDER
        c.alignment = WRAP if wrap else CENTER


def widths(ws, cols):
    for i, w in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="595959")


def score_scale(ws, ref):
    ws.conditional_formatting.add(
        ref,
        ColorScaleRule(start_type="num", start_value=1, start_color="F8696B",
                       mid_type="num", mid_value=3, mid_color="FFEB84",
                       end_type="num", end_value=5, end_color="63BE7B"),
    )


def build():
    wb = Workbook()

    # ---------------- Criteria & Weights ----------------
    cw = wb.active
    cw.title = "Criteria & Weights"
    title(cw, "Evaluation criteria and weights",
          "Yellow cells are inputs. Base weights drive the Scoring Matrix; scenario weights drive the Sensitivity sheet. Each column must sum to 100.")
    head = ["#", "Criterion", "What good looks like at Takasbank"] + [s["name"] for s in d.SCENARIOS]
    header(cw, 4, head)
    for i, c in enumerate(d.CRITERIA):
        r = 5 + i
        body(cw, r, [i + 1, c["name"], c["what"]])
        for j, s in enumerate(d.SCENARIOS):
            cell = cw.cell(row=r, column=4 + j, value=s["weights"][c["id"]])
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.alignment = CENTER
    last = 5 + len(d.CRITERIA) - 1
    tr = last + 1
    cw.cell(row=tr, column=2, value="Total (must be 100)").font = BOLD
    for j in range(len(d.SCENARIOS)):
        col = get_column_letter(4 + j)
        c = cw.cell(row=tr, column=4 + j, value=f"=SUM({col}5:{col}{last})")
        c.font = BOLD
        c.alignment = CENTER
        c.border = BORDER
    cw.conditional_formatting.add(f"D{tr}:{get_column_letter(3 + len(d.SCENARIOS))}{tr}",
                                  CellIsRule(operator="notEqual", formula=["100"], fill=FAIL_FILL))
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="40", allow_blank=False)
    dv.error = "Weight must be a whole number between 0 and 40"
    cw.add_data_validation(dv)
    dv.add(f"D5:{get_column_letter(3 + len(d.SCENARIOS))}{last}")
    widths(cw, [5, 42, 70, 16, 16, 16, 16])
    cw.freeze_panes = "D5"
    WEIGHT_ROW = {c["id"]: 5 + i for i, c in enumerate(d.CRITERIA)}

    # ---------------- Scoring Matrix ----------------
    sm = wb.create_sheet("Scoring Matrix")
    title(sm, "Weighted scoring matrix (scores 1-5, base weights)",
          "Each tool is scored as a deployable stack incl. required CDC/orchestration companions. Weighted total = SUMPRODUCT(weights, scores) / 5, so 100 = perfect.")
    header(sm, 4, ["Criterion", "Weight"] + TOOL_SHORT)
    for i, c in enumerate(d.CRITERIA):
        r = 5 + i
        body(sm, r, [c["name"]])
        wc = sm.cell(row=r, column=2, value=f"='Criteria & Weights'!D{WEIGHT_ROW[c['id']]}")
        wc.border = BORDER
        wc.alignment = CENTER
        for j, tid in enumerate(TOOL_IDS):
            cell = sm.cell(row=r, column=3 + j, value=d.SCORES[tid][c["id"]][0])
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.alignment = CENTER
    last = 5 + len(d.CRITERIA) - 1
    score_scale(sm, f"C5:{get_column_letter(2 + len(TOOL_IDS))}{last}")
    tot_row, rank_row, gate_row = last + 2, last + 3, last + 4
    sm.cell(row=tot_row, column=1, value="Weighted total (0-100)").font = BOLD
    sm.cell(row=rank_row, column=1, value="Rank").font = BOLD
    sm.cell(row=gate_row, column=1, value="Knock-out gate status (see Gates)").font = BOLD
    for j in range(len(TOOL_IDS)):
        col = get_column_letter(3 + j)
        t = sm.cell(row=tot_row, column=3 + j, value=f"=ROUND(SUMPRODUCT($B$5:$B${last},{col}5:{col}{last})/5,1)")
        t.font = BOLD
        t.alignment = CENTER
        t.border = BORDER
        rk = sm.cell(row=rank_row, column=3 + j,
                     value=f"=RANK({col}{tot_row},$C${tot_row}:${get_column_letter(2 + len(TOOL_IDS))}${tot_row})")
        rk.alignment = CENTER
        rk.border = BORDER
        g = sm.cell(row=gate_row, column=3 + j, value=f"=Gates!{get_column_letter(3 + j)}{5 + len(d.GATES)}")
        g.alignment = CENTER
        g.border = BORDER
    sm.conditional_formatting.add(f"C{tot_row}:{get_column_letter(2 + len(TOOL_IDS))}{tot_row}",
                                  ColorScaleRule(start_type="num", start_value=40, start_color="F8696B",
                                                 mid_type="num", mid_value=60, mid_color="FFEB84",
                                                 end_type="num", end_value=80, end_color="63BE7B"))
    gate_ref = f"C{gate_row}:{get_column_letter(2 + len(TOOL_IDS))}{gate_row}"
    sm.conditional_formatting.add(gate_ref, CellIsRule(operator="equal", formula=['"Fail"'], fill=FAIL_FILL))
    sm.conditional_formatting.add(gate_ref, CellIsRule(operator="equal", formula=['"Conditional"'], fill=COND_FILL))
    sm.conditional_formatting.add(gate_ref, CellIsRule(operator="equal", formula=['"Pass"'], fill=PASS_FILL))

    rat_start = gate_row + 3
    sm.cell(row=rat_start - 1, column=1, value="Scoring rationale").font = Font(bold=True, size=12, color=NAVY)
    header(sm, rat_start, ["Criterion", "Weight"] + TOOL_SHORT)
    for i, c in enumerate(d.CRITERIA):
        r = rat_start + 1 + i
        body(sm, r, [c["name"], c["weight"]] + [f"{d.SCORES[tid][c['id']][0]} - {d.SCORES[tid][c['id']][1]}" for tid in TOOL_IDS])
        sm.row_dimensions[r].height = 75
    widths(sm, [42, 9, 34, 34, 34, 34, 34])
    sm.freeze_panes = "C5"

    chart = BarChart()
    chart.type = "col"
    chart.title = "Weighted total by tool (base weights, 0-100)"
    chart.y_axis.title = "Weighted score (0-100)"
    chart.x_axis.title = "ETL tool (as deployable stack)"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 100
    chart.legend = None
    data_ref = Reference(sm, min_col=3, max_col=2 + len(TOOL_IDS), min_row=tot_row, max_row=tot_row)
    cats = Reference(sm, min_col=3, max_col=2 + len(TOOL_IDS), min_row=4, max_row=4)
    chart.add_data(data_ref, from_rows=True, titles_from_data=False)
    chart.set_categories(cats)
    chart.height, chart.width = 8, 16
    sm.add_chart(chart, f"I4")

    # ---------------- Gates ----------------
    gs = wb.create_sheet("Gates")
    title(gs, "Knock-out gates (pass / conditional / fail before scoring)",
          "A Fail removes the tool from the shortlist regardless of its weighted score. Conditional = passes only with the stated mitigation.")
    header(gs, 4, ["Gate", "Requirement"] + TOOL_SHORT)
    for i, g in enumerate(d.GATES):
        r = 5 + i
        body(gs, r, [g["id"], g["name"]] + [g["results"][tid][0] for tid in TOOL_IDS])
        for j in range(len(TOOL_IDS)):
            gs.cell(row=r, column=3 + j).alignment = CENTER
    r = 5 + len(d.GATES)
    gs.cell(row=r, column=2, value="Overall").font = BOLD
    for j in range(len(TOOL_IDS)):
        col = get_column_letter(3 + j)
        rng = f"{col}5:{col}{r - 1}"
        c = gs.cell(row=r, column=3 + j,
                    value=f'=IF(COUNTIF({rng},"Fail")>0,"Fail",IF(COUNTIF({rng},"Conditional")>0,"Conditional","Pass"))')
        c.font = BOLD
        c.alignment = CENTER
        c.border = BORDER
    ref = f"C5:{get_column_letter(2 + len(TOOL_IDS))}{r}"
    gs.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Fail"'], fill=FAIL_FILL))
    gs.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Conditional"'], fill=COND_FILL))
    gs.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Pass"'], fill=PASS_FILL))
    nr = r + 2
    gs.cell(row=nr, column=1, value="Gate evidence").font = Font(bold=True, size=12, color=NAVY)
    header(gs, nr + 1, ["Gate", "Requirement"] + TOOL_SHORT)
    for i, g in enumerate(d.GATES):
        rr = nr + 2 + i
        body(gs, rr, [g["id"], g["name"]] + [f"{g['results'][tid][0]}: {g['results'][tid][1]}" for tid in TOOL_IDS])
        gs.row_dimensions[rr].height = 60
    widths(gs, [7, 48, 28, 28, 28, 28, 28])
    gs.freeze_panes = "C5"

    # ---------------- Sensitivity ----------------
    se = wb.create_sheet("Sensitivity")
    title(se, "Sensitivity: weighted totals under alternative weightings",
          "Totals are live formulas over 'Criteria & Weights' scenario columns and 'Scoring Matrix' scores.")
    header(se, 4, ["Scenario"] + TOOL_SHORT + ["Winner"])
    first_score_row, last_score_row = 5, 5 + len(d.CRITERIA) - 1
    for i, s in enumerate(d.SCENARIOS):
        r = 5 + i
        body(se, r, [s["name"]])
        wcol = get_column_letter(4 + i)
        for j in range(len(TOOL_IDS)):
            scol = get_column_letter(3 + j)
            c = se.cell(row=r, column=2 + j,
                        value=f"=ROUND(SUMPRODUCT('Criteria & Weights'!${wcol}${first_score_row}:${wcol}${last_score_row},"
                              f"'Scoring Matrix'!{scol}${first_score_row}:{scol}${last_score_row})/5,1)")
            c.alignment = CENTER
            c.border = BORDER
        lastcol = get_column_letter(1 + len(TOOL_IDS))
        w = se.cell(row=r, column=2 + len(TOOL_IDS),
                    value=f"=INDEX($B$4:${lastcol}$4,MATCH(MAX(B{r}:{lastcol}{r}),B{r}:{lastcol}{r},0))")
        w.font = BOLD
        w.alignment = CENTER
        w.border = BORDER
    score_ref = f"B5:{get_column_letter(1 + len(TOOL_IDS))}{4 + len(d.SCENARIOS)}"
    se.conditional_formatting.add(score_ref, ColorScaleRule(start_type="num", start_value=40, start_color="F8696B",
                                                            mid_type="num", mid_value=60, mid_color="FFEB84",
                                                            end_type="num", end_value=80, end_color="63BE7B"))
    note_r = 6 + len(d.SCENARIOS)
    se.cell(row=note_r, column=1, value="Reading: ODI leads in the base, Oracle-team-first and real-time-first scenarios; dbt overtakes ODI only when TCO is weighted at 22%. PowerCenter's total is irrelevant because it fails gate G2.").alignment = WRAP
    se.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=2 + len(TOOL_IDS))
    se.row_dimensions[note_r].height = 45
    widths(se, [30, 14, 14, 14, 14, 14, 16])

    chart2 = BarChart()
    chart2.type = "col"
    chart2.grouping = "clustered"
    chart2.title = "Weighted total by scenario and tool (0-100)"
    chart2.y_axis.title = "Weighted score (0-100)"
    chart2.x_axis.title = "Weighting scenario"
    chart2.y_axis.scaling.min = 0
    chart2.y_axis.scaling.max = 100
    data2 = Reference(se, min_col=2, max_col=1 + len(TOOL_IDS), min_row=4, max_row=4 + len(d.SCENARIOS))
    cats2 = Reference(se, min_col=1, min_row=5, max_row=4 + len(d.SCENARIOS))
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.height, chart2.width = 9, 20
    se.add_chart(chart2, f"A{note_r + 2}")

    # ---------------- Capability Detail ----------------
    cd = wb.create_sheet("Capability Detail")
    title(cd, "Tool profiles", f"Vendor facts verified {d.AS_OF}; see Sources.")
    rows = [("Product / version", "version"), ("Category", "category"), ("Deployable stack evaluated", "stack"),
            ("Licensing", "licensing"), ("Lifecycle / support", "lifecycle")]
    header(cd, 4, ["Attribute"] + [t["name"] for t in d.TOOLS])
    r = 5
    for label, key in rows:
        body(cd, r, [label] + [t[key] for t in d.TOOLS])
        cd.cell(row=r, column=1).font = BOLD
        cd.row_dimensions[r].height = 90 if key in ("stack", "licensing", "lifecycle") else 35
        r += 1
    for label, key in [("Strengths", "strengths"), ("Weaknesses", "weaknesses")]:
        body(cd, r, [label] + ["\n".join(f"- {x}" for x in t[key]) for t in d.TOOLS])
        cd.cell(row=r, column=1).font = BOLD
        cd.row_dimensions[r].height = 210
        r += 1
    body(cd, r, ["Takasbank fit"] + [t["fit"] for t in d.TOOLS])
    cd.cell(row=r, column=1).font = BOLD
    cd.row_dimensions[r].height = 60
    widths(cd, [24, 42, 42, 42, 42, 42])
    cd.freeze_panes = "B5"

    # ---------------- Real-time & Latency ----------------
    rt = wb.create_sheet("Real-time & Latency")
    title(rt, "Real-time and operational reporting by latency class (S0-S4)",
          "Score 1-5 = how well the tool stack serves the class without touching prod Kale. Recommended pattern is tool-independent.")
    header(rt, 4, ["Class", "Meaning", "Takasbank example", "Recommended pattern"] + TOOL_SHORT)
    for i, lc in enumerate(d.LATENCY_CLASSES):
        r = 5 + i
        body(rt, r, [lc["id"], lc["name"], lc["example"], lc["recommended"]] + [lc["support"][tid][0] for tid in TOOL_IDS])
        for j in range(len(TOOL_IDS)):
            rt.cell(row=r, column=5 + j).alignment = CENTER
        rt.row_dimensions[r].height = 60
    score_scale(rt, f"E5:{get_column_letter(4 + len(TOOL_IDS))}{4 + len(d.LATENCY_CLASSES)}")
    nr = 6 + len(d.LATENCY_CLASSES)
    rt.cell(row=nr, column=1, value="Notes per tool").font = Font(bold=True, size=12, color=NAVY)
    header(rt, nr + 1, ["Class", "", "", ""] + TOOL_SHORT)
    for i, lc in enumerate(d.LATENCY_CLASSES):
        r = nr + 2 + i
        body(rt, r, [lc["id"], "", "", ""] + [lc["support"][tid][1] for tid in TOOL_IDS])
    widths(rt, [7, 26, 36, 46, 18, 18, 18, 18, 18])
    rt.freeze_panes = "E5"

    # ---------------- Use-Case Fit ----------------
    uc = wb.create_sheet("Use-Case Fit")
    title(uc, "Fit to Takasbank workloads (1-5)",
          "Workloads taken from the July 2025 DWH presentation and UG meeting notes.")
    header(uc, 4, ["ID", "Use case", "Volume / behaviour"] + TOOL_SHORT)
    for i, u in enumerate(d.USE_CASES):
        r = 5 + i
        body(uc, r, [u["id"], u["name"], u["detail"]] + [u["scores"][tid] for tid in TOOL_IDS])
        for j in range(len(TOOL_IDS)):
            c = uc.cell(row=r, column=4 + j)
            c.alignment = CENTER
            c.fill = INPUT_FILL
        uc.row_dimensions[r].height = 48
    last = 4 + len(d.USE_CASES)
    score_scale(uc, f"D5:{get_column_letter(3 + len(TOOL_IDS))}{last}")
    avg = last + 1
    uc.cell(row=avg, column=2, value="Average fit").font = BOLD
    for j in range(len(TOOL_IDS)):
        col = get_column_letter(4 + j)
        c = uc.cell(row=avg, column=4 + j, value=f"=ROUND(AVERAGE({col}5:{col}{last}),2)")
        c.font = BOLD
        c.alignment = CENTER
        c.border = BORDER
    widths(uc, [6, 30, 70, 13, 13, 13, 13, 13])
    uc.freeze_panes = "D5"

    # ---------------- TCO Drivers ----------------
    tc = wb.create_sheet("TCO Drivers")
    title(tc, "5-year TCO drivers (bands, not quotes)",
          "List prices are public reference points only; Takasbank must obtain quotes. Higher band = higher cost.")
    header(tc, 4, ["Driver"] + TOOL_SHORT)
    for i, row in enumerate(d.TCO):
        r = 5 + i
        body(tc, r, [row["driver"]] + [row[tid] for tid in TOOL_IDS])
        tc.cell(row=r, column=1).font = BOLD
        tc.row_dimensions[r].height = 55
    band_row = 4 + len(d.TCO)
    band_ref = f"B{band_row}:{get_column_letter(1 + len(TOOL_IDS))}{band_row}"
    tc.conditional_formatting.add(band_ref, CellIsRule(operator="equal", formula=['"Low"'], fill=PASS_FILL))
    tc.conditional_formatting.add(band_ref, CellIsRule(operator="equal", formula=['"Medium"'], fill=COND_FILL))
    tc.conditional_formatting.add(band_ref, CellIsRule(operator="equal", formula=['"High"'], fill=FAIL_FILL))
    tc.conditional_formatting.add(band_ref, CellIsRule(operator="equal", formula=['"Very high"'], fill=FAIL_FILL))
    widths(tc, [28, 36, 36, 36, 36, 36])
    tc.freeze_panes = "B5"

    # ---------------- Risks ----------------
    rk = wb.create_sheet("Risks")
    title(rk, "Risk register by tool")
    header(rk, 4, ["Tool", "Risk", "Likelihood", "Impact", "Mitigation"])
    for i, x in enumerate(d.RISKS):
        r = 5 + i
        body(rk, r, [x["tool"], x["risk"], x["likelihood"], x["impact"], x["mitigation"]])
        rk.cell(row=r, column=3).alignment = CENTER
        rk.cell(row=r, column=4).alignment = CENTER
        rk.row_dimensions[r].height = 45
    ref = f"C5:D{4 + len(d.RISKS)}"
    rk.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"High"'], fill=FAIL_FILL))
    rk.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Medium"'], fill=COND_FILL))
    rk.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Low"'], fill=PASS_FILL))
    rk.auto_filter.ref = f"A4:E{4 + len(d.RISKS)}"
    widths(rk, [13, 52, 12, 12, 72])
    rk.freeze_panes = "A5"

    # ---------------- PoC Plan ----------------
    pp = wb.create_sheet("PoC Plan")
    title(pp, "Proof of concept: ODI vs dbt head-to-head on Takasbank data",
          "Run on the Kale standby and DSS, never prod Kale. PowerCenter excluded by gate G2; CDP/Spark excluded by fit.")
    header(pp, 4, ["Test", "Name", "Use case", "What is measured", "Acceptance criterion", "ODI result", "dbt result"])
    for i, t in enumerate(d.POC_TESTS):
        r = 5 + i
        body(pp, r, [t["id"], t["name"], t["use_case"], t["measure"], t["accept"], "", ""])
        pp.cell(row=r, column=6).fill = INPUT_FILL
        pp.cell(row=r, column=7).fill = INPUT_FILL
        pp.row_dimensions[r].height = 48
    dv2 = DataValidation(type="list", formula1='"Pass,Partial,Fail,Not run"', allow_blank=True)
    pp.add_data_validation(dv2)
    dv2.add(f"F5:G{4 + len(d.POC_TESTS)}")
    widths(pp, [6, 28, 10, 55, 55, 12, 12])
    pp.freeze_panes = "A5"

    # ---------------- Context & Requirements ----------------
    cx = wb.create_sheet("Context & Requirements")
    title(cx, "Takasbank context that drives the evaluation",
          "Facts from the Obsidian vault notes (meetings, presentations, topology). Not vendor claims.")
    header(cx, 4, ["Topic", "Fact"])
    for i, (k, v) in enumerate(d.CONTEXT):
        body(cx, 5 + i, [k, v])
        cx.cell(row=5 + i, column=1).font = BOLD
        cx.row_dimensions[5 + i].height = 35
    widths(cx, [20, 120])

    # ---------------- Sources ----------------
    so = wb.create_sheet("Sources")
    title(so, f"Sources (verified {d.AS_OF})")
    header(so, 4, ["#", "Source", "URL"])
    for i, (name, url) in enumerate(d.SOURCES):
        r = 5 + i
        body(so, r, [i + 1, name, url])
        so.cell(row=r, column=3).hyperlink = url
        so.cell(row=r, column=3).font = Font(color="0563C1", underline="single")
    widths(so, [5, 60, 110])

    # ---------------- Summary (first sheet) ----------------
    su = wb.create_sheet("Summary", 0)
    title(su, "Takasbank on-prem EDW - ETL tool comparison",
          f"ODI, Informatica PowerCenter, dbt, Cloudera Data Platform, Apache Spark  |  as of {d.AS_OF}")
    su["A4"] = "Recommendation"
    su["A4"].font = Font(bold=True, size=12, color=NAVY)
    rec = [
        "1. Select Oracle Data Integrator 14c as the enterprise ELT platform: pushdown into the existing Exadata, PL/SQL-friendly, supported to 2030/2033.",
        "2. Serve near-real-time operational ('izleme') reports from the Kale Active Data Guard read-only standby, not from an ETL tool.",
        "3. Use GoldenGate (via ODI JKM) only for data products whose S0 SLA is proven, after the licence question is settled.",
        "4. Keep dbt-core + dbt-oracle as the qualified alternative and procurement-free fallback; run the PoC head-to-head ODI vs dbt.",
        "5. Exclude PowerCenter (fails support-horizon gate; successor needs a cloud control plane). Do not build the core DWH on Cloudera or Spark.",
    ]
    for i, line in enumerate(rec):
        su.cell(row=5 + i, column=1, value=line).alignment = WRAP
        su.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=7)
        su.row_dimensions[5 + i].height = 30
    hr = 11
    header(su, hr, ["Metric"] + TOOL_SHORT)
    lines = [
        ("Weighted total (base)", lambda j: f"='Scoring Matrix'!{get_column_letter(3 + j)}{5 + len(d.CRITERIA) + 1}"),
        ("Rank (base)", lambda j: f"='Scoring Matrix'!{get_column_letter(3 + j)}{5 + len(d.CRITERIA) + 2}"),
        ("Knock-out gates", lambda j: f"=Gates!{get_column_letter(3 + j)}{5 + len(d.GATES)}"),
        ("Average use-case fit (1-5)", lambda j: f"='Use-Case Fit'!{get_column_letter(4 + j)}{5 + len(d.USE_CASES)}"),
        ("5-year TCO band", lambda j: f"='TCO Drivers'!{get_column_letter(2 + j)}{4 + len(d.TCO)}"),
    ]
    for i, (label, fn) in enumerate(lines):
        r = hr + 1 + i
        body(su, r, [label])
        su.cell(row=r, column=1).font = BOLD
        for j in range(len(TOOL_IDS)):
            c = su.cell(row=r, column=2 + j, value=fn(j))
            c.alignment = CENTER
            c.border = BORDER
    for i, s in enumerate(d.SCENARIOS[1:], start=len(lines)):
        r = hr + 1 + i
        body(su, r, [f"Weighted total ({s['name']})"])
        su.cell(row=r, column=1).font = BOLD
        for j in range(len(TOOL_IDS)):
            c = su.cell(row=r, column=2 + j, value=f"=Sensitivity!{get_column_letter(2 + j)}{5 + d.SCENARIOS.index(s)}")
            c.alignment = CENTER
            c.border = BORDER
    end = hr + len(lines) + len(d.SCENARIOS) - 1
    gate_r = hr + 3
    ref = f"B{gate_r}:{get_column_letter(1 + len(TOOL_IDS))}{gate_r}"
    su.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Fail"'], fill=FAIL_FILL))
    su.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Conditional"'], fill=COND_FILL))
    su.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Pass"'], fill=PASS_FILL))
    su.conditional_formatting.add(f"B{hr + 1}:{get_column_letter(1 + len(TOOL_IDS))}{hr + 1}",
                                  ColorScaleRule(start_type="num", start_value=40, start_color="F8696B",
                                                 mid_type="num", mid_value=60, mid_color="FFEB84",
                                                 end_type="num", end_value=80, end_color="63BE7B"))
    nr = end + 2
    su.cell(row=nr, column=1, value="How to use this workbook").font = Font(bold=True, size=12, color=NAVY)
    guide = [
        "Yellow cells are inputs: weights on 'Criteria & Weights', scores on 'Scoring Matrix' and 'Use-Case Fit', PoC results on 'PoC Plan'.",
        "All totals, ranks, gate outcomes and scenario winners are formulas and recalculate on change.",
        "The tools are not like-for-like (suite vs framework vs platform vs engine); each is scored as a deployable stack with its CDC/orchestration companions.",
        "Vault note: Obsidian Vault/Kendi Notlarım/ETL Araç Karşılaştırması - Ekim 2026.md. Source of numbers: comparison_data.py.",
    ]
    for i, g in enumerate(guide):
        su.cell(row=nr + 1 + i, column=1, value=f"- {g}").alignment = WRAP
        su.merge_cells(start_row=nr + 1 + i, start_column=1, end_row=nr + 1 + i, end_column=7)
        su.row_dimensions[nr + 1 + i].height = 30
    widths(su, [34, 16, 16, 16, 16, 16, 4])

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(OUT)
    return OUT


def export_json():
    payload = {
        "asOf": d.AS_OF,
        "tools": d.TOOLS,
        "criteria": d.CRITERIA,
        "scores": {tid: {cid: {"score": s, "why": w} for cid, (s, w) in crit.items()} for tid, crit in d.SCORES.items()},
        "gates": [{"id": g["id"], "name": g["name"],
                   "results": {tid: {"status": s, "note": n} for tid, (s, n) in g["results"].items()}} for g in d.GATES],
        "scenarios": d.SCENARIOS,
        "latency": [{**{k: v for k, v in lc.items() if k != "support"},
                     "support": {tid: {"score": s, "note": n} for tid, (s, n) in lc["support"].items()}}
                    for lc in d.LATENCY_CLASSES],
        "useCases": d.USE_CASES,
        "tco": d.TCO,
        "risks": d.RISKS,
        "poc": d.POC_TESTS,
        "totals": {s["id"]: {tid: d.weighted_total(tid, s["weights"]) for tid in TOOL_IDS} for s in d.SCENARIOS},
    }
    path = HERE / "comparison-data.json"
    path.write_text(json.dumps(payload, indent=1, ensure_ascii=False))
    return path


if __name__ == "__main__":
    print(build())
    if "--json" in sys.argv:
        print(export_json())
